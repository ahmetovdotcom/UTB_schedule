"""Export the current timetable: one course per workbook, departments as sheets."""
from collections import defaultdict
from io import BytesIO
from copy import copy, deepcopy
from pathlib import Path
from uuid import uuid4
import re

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload
from starlette.concurrency import run_in_threadpool

from database.models import AcademicLoad, Room, ScheduleItem, StudentGroup, TimeSlot, get_session
from room_rules import LESSON_TYPE_LABELS

router = APIRouter()
DAYS = ('Понедельник', 'Вторник', 'Среда', 'Четверг', 'Пятница', 'Суббота')
XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def sheet_name(name, used):
    base = re.sub(r'[\\/*?:\[\]]', ' ', ILLEGAL_CHARACTERS_RE.sub('', name)).strip().strip("'")[:31] or 'Кафедра'
    # Excel reserves History and compares worksheet titles case-insensitively.
    if base.casefold() == 'history':
        base = 'History кафедра'
    result, suffix = base, 2
    while result.casefold() in used:
        tail = f' ({suffix})'
        result = base[:31 - len(tail)] + tail
        suffix += 1
    used.add(result.casefold())
    return result


def text_cell(sheet, row, col, text):
    cell = sheet.cell(row, col)
    cell.value = ILLEGAL_CHARACTERS_RE.sub('', str(text))
    # User-entered names must stay literal even when they begin with '='.
    cell.data_type = 's'
    return cell


TEMPLATE_PATH = Path(__file__).with_name('templates') / 'timetable.xlsx'
# Original day heading, first time row and last time row, including its blank row.
TEMPLATE_DAYS = ((13, 14, 23), (24, 25, 34), (35, 36, 45), (46, 47, 57), (58, 59, 69))


def time_key(start, end=None):
    if end is None:
        start, end = start.replace('–', '-').split('-', 1)
    return (start.strip().replace('.', ':'), end.strip().replace('.', ':'))


def build_workbook(course, groups, slots, lessons):
    workbook = load_workbook(TEMPLATE_PATH)
    template = workbook.worksheets[0]
    template.title = uuid4().hex[:31]
    departments = defaultdict(list)
    for group in groups:
        departments[(group['department_id'], group['department_name'] or 'Без кафедры')].append(group)
    slot_by_id = {slot['id']: time_key(slot['start'], slot['end']) for slot in slots}
    by_cell = defaultdict(list)
    for lesson in lessons:
        for group_id in lesson['group_ids']:
            by_cell[(group_id, lesson['day'], slot_by_id[lesson['slot_id']])].append(lesson)
    used = set()
    for (_department_id, name), members in sorted(
        departments.items(), key=lambda entry: (entry[0][0] is None, entry[0][1].casefold(), entry[0][0] or 0),
    ):
        sheet = workbook.create_sheet(sheet_name(name, used))
        members.sort(key=lambda g: (g['code'].casefold(), g['id']))
        # Keep the original eight group blocks, including empty ones. Larger
        # departments extend the same four-column pattern on the same sheet.
        last_col = max(33, 1 + 4 * len(members))
        if last_col > 16384:
            raise ValueError('Слишком много групп одной кафедры для листа Excel')
        member_ids = {g['id'] for g in members}
        relevant = [lesson for lesson in lessons if member_ids.intersection(lesson['group_ids'])]
        # (source row, day, time, original row). Keep every original empty row;
        # only extend the template when scheduled times/days require it.
        plan = [(row, None, None, True) for row in range(1, 13)]
        days = 6 if any(lesson['day'] == 6 for lesson in relevant) else 5
        for day in range(1, days + 1):
            if day <= 5:
                heading, first, last = TEMPLATE_DAYS[day - 1]
                plan.append((heading, day, None, True))
                time_rows, blanks = [], []
                for row in range(first, last + 1):
                    value = template.cell(row, 1).value
                    if value:
                        time_rows.append((row, day, time_key(value), True))
                    else:
                        blanks.append((row, day, None, True))
            else:
                plan.append((13, day, None, False))
                time_rows, blanks = [], []
            existing = {entry[2] for entry in time_rows}
            required = {slot_by_id[l['slot_id']] for l in relevant if l['day'] == day}
            for key in sorted(required - existing):
                time_rows.append((14, day, key, False))
            plan.extend(sorted(time_rows, key=lambda entry: entry[2]))
            plan.extend(blanks)
        plan.extend((row, None, None, True) for row in range(70, 83))
        origin_rows = {}
        for row, (source_row, day, key, original) in enumerate(plan, 1):
            if original:
                origin_rows[source_row] = row
            dimension = copy(template.row_dimensions[source_row])
            dimension.index = row
            sheet.row_dimensions[row] = dimension
            for col in range(1, max(template.max_column, last_col) + 1):
                source_col = col if col <= 33 or col > last_col else 2 + (col - 34) % 4
                source = template.cell(source_row, source_col)
                dest = sheet.cell(row, col)
                dest._style = copy(source._style)
                if col <= 33 and not isinstance(source, MergedCell):
                    dest.value = source.value
                    dest.data_type = source.data_type
            if day and key:
                text_cell(sheet, row, 1, '-'.join(part.replace(':', '.') for part in key))
                for index, group in enumerate(members):
                    col = 2 + 4 * index
                    entries = by_cell[(group['id'], day, key)]
                    descriptions, rooms = [], []
                    for lesson in entries:
                        short_type = {'Лекция': 'лек.', 'Практика': 'пр.', 'Лабораторная': 'лаб.'}.get(lesson['type'], lesson['type'])
                        descriptions.append(f"{lesson['subject']}-{short_type} {lesson['teacher']}")
                        rooms.append(lesson['room'])
                    text_cell(sheet, row, col, '\n'.join(descriptions))
                    text_cell(sheet, row, col + 3, '\n'.join(rooms))
            elif day and not key and (source_row in (13, 24, 35, 46, 58)):
                text_cell(sheet, row, 1, DAYS[day - 1].upper())
            if not original:
                if key:
                    for col in range(2, last_col + 1, 4):
                        sheet.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 2)
                else:
                    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
        for merged in template.merged_cells.ranges:
            if merged.min_row not in origin_rows or merged.max_row not in origin_rows:
                continue
            # Some rooms in the example span several lesson rows. Preserve
            # those merges only if all new values match; otherwise a later
            # lesson's room would be silently discarded by merge_cells.
            if merged.min_row >= 14 and merged.max_row > merged.min_row:
                values = {sheet.cell(row, merged.min_col).value for row in range(
                    origin_rows[merged.min_row], origin_rows[merged.max_row] + 1,
                )}
                if len(values) > 1:
                    for row in range(origin_rows[merged.min_row], origin_rows[merged.max_row] + 1):
                        sheet.cell(row, merged.min_col)._style = copy(template.cell(merged.min_row, merged.min_col)._style)
                    continue
            right = merged.max_col
            if merged.min_col == 1 and merged.max_col == 33:
                right = last_col
            elif merged.min_row <= 5 and merged.max_col == 32:
                right = last_col - 1
            sheet.merge_cells(start_row=origin_rows[merged.min_row], start_column=merged.min_col,
                              end_row=origin_rows[merged.max_row], end_column=right)
        for key, dimension in template.column_dimensions.items():
            sheet.column_dimensions[key] = copy(dimension)
        for col in range(34, last_col + 1):
            # First group B:E has one shared width in the original file.
            sheet.column_dimensions[get_column_letter(col)].width = template.column_dimensions['B'].width
        # Additional groups need the same header and three-column lesson merge.
        for col in range(34, last_col + 1, 4):
            sheet.merge_cells(start_row=11, start_column=col, end_row=11, end_column=col + 3)
            sheet.merge_cells(start_row=12, start_column=col, end_row=12, end_column=col + 2)
            text_cell(sheet, 12, col, template['B12'].value)
            text_cell(sheet, 12, col + 3, template['E12'].value)
            for row, (source_row, day, key, original) in enumerate(plan, 1):
                if original and day and source_row not in (13, 24, 35, 46, 58):
                    sheet.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 2)
        for index, group in enumerate(members):
            text_cell(sheet, 11, 2 + 4 * index, f"{group['code']} ({group['students_count']})")
        text_cell(sheet, 4, 6, re.sub(r'^\d+ курс', f'{course} курс', template['F4'].value))
        text_cell(sheet, 5, 6, re.sub(r'^\d+ семестр', f'{course * 2 - 1} семестр', template['F5'].value))
        sheet.sheet_format = copy(template.sheet_format)
        sheet.sheet_properties = copy(template.sheet_properties)
        sheet.views = deepcopy(template.views)
        sheet.page_setup = copy(template.page_setup)
        sheet.page_margins = copy(template.page_margins)
        sheet.print_options = copy(template.print_options)
        sheet.HeaderFooter = deepcopy(template.HeaderFooter)
        sheet.row_breaks = deepcopy(template.row_breaks)
        sheet.col_breaks = deepcopy(template.col_breaks)
        if template.print_area:
            sheet.print_area = template.print_area
        sheet.print_title_rows = template.print_title_rows
        sheet.print_title_cols = template.print_title_cols
    workbook.remove(template)
    workbook.active = 0
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@router.get('/api/schedule/export.xlsx')
async def export_schedule(course: int = Query(ge=1), session: AsyncSession = Depends(get_session)):
    groups = (await session.scalars(select(StudentGroup).where(StudentGroup.course == course)
        .options(selectinload(StudentGroup.department)))).all()
    if not groups:
        raise HTTPException(404, 'Для выбранного курса нет групп')
    slots = (await session.scalars(select(TimeSlot).order_by(TimeSlot.slot_number, TimeSlot.id))).all()
    items = (await session.scalars(select(ScheduleItem).where(
        ScheduleItem.academic_load.has(AcademicLoad.groups.any(StudentGroup.course == course))
    ).options(
        selectinload(ScheduleItem.academic_load).selectinload(AcademicLoad.groups),
        selectinload(ScheduleItem.academic_load).selectinload(AcademicLoad.subject),
        selectinload(ScheduleItem.academic_load).selectinload(AcademicLoad.teacher),
        selectinload(ScheduleItem.room).selectinload(Room.building),
    ).order_by(ScheduleItem.id))).all()
    group_data = [dict(id=g.id, code=g.code, students_count=g.students_count,
                       department_id=g.department_id, department_name=g.department.name if g.department else None) for g in groups]
    slot_data = [dict(id=s.id, number=s.slot_number, start=s.start_time, end=s.end_time) for s in slots]
    lessons = []
    for item in items:
        load = item.academic_load
        lessons.append(dict(day=item.day_of_week, slot_id=item.time_slot_id,
            group_ids=[g.id for g in load.groups], group_codes=[g.code for g in sorted(load.groups, key=lambda g: (g.code, g.id))],
            subject=load.subject.name, type=LESSON_TYPE_LABELS.get(load.lesson_type, load.lesson_type),
            teacher=load.teacher.full_name, room=f'{item.room.building.name}/{item.room.number}'))
    try:
        content = await run_in_threadpool(build_workbook, course, group_data, slot_data, lessons)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    return Response(content, media_type=XLSX_TYPE, headers={
        'Content-Disposition': f'attachment; filename="timetable-course-{course}.xlsx"',
        'Cache-Control': 'no-store',
    })
