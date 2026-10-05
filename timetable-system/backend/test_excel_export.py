"""Excel export, group departments and additive migration regression checks."""
import asyncio
from io import BytesIO
from copy import copy
import sqlite3
import httpx
from openpyxl import load_workbook
from sqlalchemy import select
from auth import require_user
from database.group_department_migration import migrate_group_department
from database.models import Department, StudentGroup, get_session
from excel_export import build_workbook, sheet_name
from main import app, get_initial_data, AssignRequest, assign_schedule
from references import GroupIn, create_group, update_group, delete_department
from test_validation import make_session, seed_minimal


def check_migration():
    db = sqlite3.connect(':memory:')
    db.executescript('CREATE TABLE departments (id INTEGER PRIMARY KEY, name TEXT); '
                     'CREATE TABLE student_groups (id INTEGER PRIMARY KEY, code TEXT, course INTEGER, students_count INTEGER, publication_key TEXT);')
    db.executemany('INSERT INTO student_groups VALUES (?, ?, 1, 20, ?)',
                   [(i, f'G{i}', f'key{i}') for i in range(10000)])
    execute = lambda sql: db.execute(sql).fetchall()
    migrate_group_department(execute)
    migrate_group_department(execute)
    assert execute('SELECT count(*) FROM student_groups WHERE department_id IS NULL')[0][0] == 10000
    assert execute('SELECT code, publication_key FROM student_groups WHERE id = 9999') == [('G9999', 'key9999')]
    db.close()
    print('OK: migration preserves 10000 groups and is idempotent')


def check_names_and_text():
    used = set()
    names = [sheet_name(n, used) for n in ['a/b', 'a:b', 'A:B', 'x' * 40, 'x' * 41, "'", 'History', 'Без кафедры', 'Без кафедры']]
    assert len(set(n.casefold() for n in names)) == len(names)
    assert all(0 < len(n) <= 31 for n in names)
    groups = [dict(id=1, code='=2+2', students_count=20, department_id=1, department_name='=SUM(1,2)')]
    slots = [dict(id=1, number=1, start='08:00', end='08:50')]
    lessons = [dict(group_ids=[1], group_codes=['=2+2'], day=1, slot_id=1,
                    subject='=HYPERLINK("https://example.invalid")', type='Лекция', teacher='Teacher', room='Room')]
    book = load_workbook(BytesIO(build_workbook(1, groups, slots, lessons)))
    assert book.active['B14'].data_type == 's' and book.active['B14'].value.startswith('=HYPERLINK')
    assert all(c.data_type != 'f' for s in book for row in s for c in row)
    from excel_export import TEMPLATE_PATH
    template = load_workbook(TEMPLATE_PATH).worksheets[0]
    sheet = book.active
    assert {str(m) for m in sheet.merged_cells.ranges} == {str(m) for m in template.merged_cells.ranges}
    for coordinate in ['G1', 'G3', 'F4', 'F5', 'AC6', 'AC8', 'B11', 'B12', 'A13', 'A14', 'B14', 'E14', 'R72']:
        actual, expected = sheet[coordinate], template[coordinate]
        assert copy(actual.font) == copy(expected.font), coordinate
        assert copy(actual.alignment) == copy(expected.alignment), coordinate
        assert copy(actual.border) == copy(expected.border), coordinate
        assert copy(actual.fill) == copy(expected.fill), coordinate
    assert sheet.page_setup == template.page_setup
    assert sheet.page_margins == template.page_margins
    assert sheet['G1'].value == template['G1'].value
    assert sheet['AC8'].value == template['AC8'].value
    assert sheet['F4'].value.startswith('1 курс')
    assert sheet['F5'].value.startswith('1 семестр')
    for row in range(1, 83):
        assert sheet.row_dimensions[row].height == template.row_dimensions[row].height
    for key, dimension in template.column_dimensions.items():
        assert sheet.column_dimensions[key].width == dimension.width
    assert sheet['F11'].value is None
    print('OK: exact template styles, merges, dimensions, headers and printing; user text stays literal')


def check_template_extensions():
    groups = [dict(id=i, code=f'G{i:02}', students_count=20, department_id=1, department_name='ТФ 3 курс типо') for i in range(1, 11)]
    slots = [dict(id=1, number=1, start='08:30', end='09:20'),
             dict(id=2, number=2, start='14:00', end='14:50')]
    lessons = [dict(group_ids=[10], group_codes=['G10'], day=6, slot_id=1,
                    subject='Субботняя пара', type='Лекция', teacher='Teacher', room='1/999'),
               dict(group_ids=[4], group_codes=['G04'], day=2, slot_id=2,
                    subject='Внутри старого объединения', type='Лекция', teacher='Teacher', room='1/123')]
    book = load_workbook(BytesIO(build_workbook(3, groups, slots, lessons)))
    sheet = book.active
    assert sheet.title == 'ТФ 3 курс типо'
    values = [c.value for row in sheet for c in row if c.value]
    assert 'СУББОТА' in values and '08.30-09.20' in values
    assert 'G10 (20)' in values and '1/999' in values and '1/123' in values
    assert any('Субботняя пара' in str(value) for value in values)
    assert sheet['Q31'].value == '1/123'
    assert 'Q30:Q34' not in {str(m) for m in sheet.merged_cells.ranges}
    print('OK: extra groups, Saturday, custom times and original merged-room cells retain all lessons')


async def check_api():
    session, engine = await make_session()
    try:
        data = await seed_minimal(session)
        departments = [Department(name='Информационные технологии'), Department(name='Экономика')]
        session.add_all(departments)
        await session.flush()
        groups = (await session.scalars(select(StudentGroup).order_by(StudentGroup.id))).all()
        groups[0].course = groups[1].course = 3
        groups[0].department_id = departments[0].id
        groups[1].department_id = departments[1].id
        groups[2].course = 2
        await session.commit()
        await assign_schedule(AssignRequest(academic_load_id=data['lab_load'].id, day_of_week=2,
                              time_slot_id=data['slot_1'].id, room_id=data['lab_room'].id), session)
        new_group = await create_group(GroupIn(code='Без занятий', students_count=10, course=3), session)
        await update_group(new_group['id'], GroupIn(code='Без занятий', students_count=10, course=3,
                           department_id=departments[0].id), session)
        initial = await get_initial_data(session)
        assert next(g for g in initial['groups'] if g['id'] == new_group['id'])['department_name'] == departments[0].name
        await create_group(GroupIn(code='Без кафедры группа', students_count=10, course=3), session)
        async def session_override():
            yield session
        app.dependency_overrides[get_session] = session_override
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url='http://test') as client:
            assert (await client.get('/api/schedule/export.xlsx?course=3')).status_code == 401
            app.dependency_overrides[require_user] = lambda: 'test-user'
            response = await client.get('/api/schedule/export.xlsx?course=3')
            assert response.status_code == 200, response.text
            assert response.headers['content-disposition'].endswith('timetable-course-3.xlsx"')
            book = load_workbook(BytesIO(response.content))
            assert set(book.sheetnames) == {'Информационные технологии', 'Экономика', 'Без кафедры'}
            for name in ('Информационные технологии', 'Экономика'):
                contents = '\n'.join(str(c.value) for row in book[name] for c in row if c.value)
                assert 'Предмет-лаб. Преп. Б' in contents
                assert 'Корпус/A-2' in contents
                assert groups[2].code not in contents
            assert 'Без занятий' in '\n'.join(str(c.value) for row in book['Информационные технологии'] for c in row)
            other = await client.get('/api/schedule/export.xlsx?course=2')
            assert len(load_workbook(BytesIO(other.content)).sheetnames) == 1
            assert (await client.get('/api/schedule/export.xlsx?course=99')).status_code == 404
            assert (await client.get('/api/schedule/export.xlsx?course=0')).status_code == 422
            assert (await client.get('/api/schedule/export.xlsx')).status_code == 422
            bad = await client.post('/api/groups', json=dict(code='bad', students_count=1, course=3, department_id=-1))
            assert bad.status_code == 404
        await delete_department(departments[0].id, session)
        session.expunge_all()
        assert (await session.get(StudentGroup, new_group['id'])).department_id is None
        print('OK: authenticated XLSX, courses, departments, empty groups, streams and department removal')
    finally:
        app.dependency_overrides.clear()
        await session.close()
        await engine.dispose()


if __name__ == '__main__':
    check_migration()
    check_names_and_text()
    check_template_extensions()
    asyncio.run(check_api())
